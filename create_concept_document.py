import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Sets internal padding for a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_concept_document(output_path):
    doc = docx.Document()

    # Precise 2-page fit layout (0.55 inch top/bottom, 0.65 inch sides)
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Human Professional Color Palette
    PRIMARY = RGBColor(15, 23, 42)      # Deep Navy (#0f172a)
    ACCENT = RGBColor(14, 116, 144)     # Deep Teal (#0e7490)
    BODY_COLOR = RGBColor(51, 65, 85)   # Slate 700 (#334155)
    MUTED = RGBColor(100, 116, 139)     # Slate 500 (#64748b)

    def add_para(text="", font_size=9.5, bold=False, color=BODY_COLOR, space_before=0, space_after=3.5, line_spacing=1.12):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.color.rgb = color
        return p

    def add_heading(text, space_before=6, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        return p

    # -------------------------------------------------------------
    # DOCUMENT HEADER
    # -------------------------------------------------------------
    add_para("PROJECT CONCEPT PROPOSAL  |  PUBLIC SAFETY & INTELLIGENT TRANSPORTATION", font_size=8, bold=True, color=ACCENT, space_after=1)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(1.5)
    r_title = p_title.add_run("RAKSHAK-AI: Autonomous Emergency Guardianship & Inter-Agency Emergency Resource Planning (ERP)")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY

    add_para("Author: Stevin Joseph B  |  Institution: Kristu Jayanti College (Autonomous), Bengaluru  |  Field: Cybersecurity & Distributed Systems", font_size=8.5, bold=False, color=MUTED, space_after=5)

    # Top Summary Box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.2)
    cell = table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    p_box = cell.paragraphs[0]
    p_box.paragraph_format.space_before = Pt(1)
    p_box.paragraph_format.space_after = Pt(1)
    p_box.paragraph_format.line_spacing = 1.12
    r_b1 = p_box.add_run("Concept Summary: ")
    r_b1.font.name = "Calibri"
    r_b1.font.size = Pt(9)
    r_b1.font.bold = True
    r_b1.font.color.rgb = PRIMARY
    r_b2 = p_box.add_run(
        "Rakshak-AI is an end-to-end emergency orchestration platform designed to eliminate the critical 40-minute 'Golden Hour' delay "
        "in Indian road crashes. By uniting on-device smartphone edge AI with municipal traffic signal preemption (C-V2X) and "
        "Ayushman Bharat (ABHA) medical vaults, the system coordinates victim detection, police dispatch, ambulance navigation, "
        "and hospital trauma bay preparation into a single, synchronized emergency workflow."
    )
    r_b2.font.name = "Calibri"
    r_b2.font.size = Pt(9)
    r_b2.font.color.rgb = BODY_COLOR

    # -------------------------------------------------------------
    # SECTION 1: THE GROUND REALITY
    # -------------------------------------------------------------
    add_heading("1. The Ground Reality: Why I Chose to Build This", space_before=7)
    
    add_para(
        "During my daily commutes through Bengaluru—whether along Hennur Road, through Kothanur, or past the Outer Ring Road—"
        "I have repeatedly seen the exact same tragedy. A two-wheeler skids or is hit. The rider is thrown to the road, knocked unconscious. "
        "Bystanders gather, but hesitate to help out of fear of police questioning or hospital liability. It routinely takes 15 to 20 minutes before anyone dials 108. "
        "When an ambulance finally arrives, it gets trapped in bumper-to-bumper traffic at intersections like Silk Board, blaring its siren helplessly while the light stays red. "
        "When the victim finally reaches the casualty ward, doctors lose critical minutes scrambling to determine their blood group and emergency medical history."
    )
    
    add_para(
        "Medical science defines the first 60 minutes after a crash as the 'Golden Hour'. In India, where over 1.53 lakh people die on roads every year "
        "(11% of global road deaths with only 1% of the world's vehicles), trauma surgeons state that over 50% of these fatalities are preventable "
        "if surgical intervention begins within that first hour. The problem is not medical skill; our doctors are outstanding. The failure is resource fragmentation: "
        "the victim's phone, the police, the ambulance, and the hospital operate in complete silos. I built Rakshak-AI to unite them into one synchronized system."
    )

    # -------------------------------------------------------------
    # SECTION 2: THE CORE CONCEPT & ERP ARCHITECTURE
    # -------------------------------------------------------------
    add_heading("2. The Core Concept: Emergency Resource Planning (ERP)")
    
    add_para(
        "I designed Rakshak-AI around the principles of Enterprise Resource Planning (ERP). In commercial enterprises, an ERP (like SAP) coordinates "
        "warehouse inventory, procurement, and logistics to eliminate bottlenecks. In emergency response, saving a life is an identical resource-matching problem: "
        "the demand (crash severity, location, victim profile) must be matched instantly with critical medical supply (an available ALS ambulance, "
        "a cleared traffic corridor, an open Trauma ICU bed, and matching blood units). Rakshak-AI acts as an Emergency Operations ERP, enforcing a strict "
        "finite state machine: Suspected Impact ➔ Promoted Incident ➔ Hospital Sector Matching ➔ Atomic Ambulance Claim ➔ Green Wave Transit ➔ Trauma Bay Handover."
    )

    # -------------------------------------------------------------
    # SECTION 3: HOW THE SYSTEM OPERATES STEP-BY-STEP
    # -------------------------------------------------------------
    add_heading("3. How the System Operates (From Impact to Hospital Docking)")

    # Step 1
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(1)
    p_s1.paragraph_format.space_after = Pt(2.5)
    p_s1.paragraph_format.line_spacing = 1.12
    r_h1 = p_s1.add_run("• Step 1: On-Device Edge Crash AI & Physics Filtering: ")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(9.2)
    r_h1.font.bold = True
    r_h1.font.color.rgb = PRIMARY
    r_t1 = p_s1.add_run(
        "The mobile app runs an on-device TensorFlow Lite 1D-CNN (42 KB footprint) sampling accelerometer and gyroscope sensors at 100 Hz. "
        "To eliminate false alarms from dropped phones or potholes, the edge engine inspects free-fall (a dropped phone experiences near-0G immediately before impact; "
        "a vehicular crash shows no pre-impact free-fall), GPS velocity drop, and post-impact stillness. For two-wheelers, the algorithm bypasses airbag requirements "
        "and monitors angular tumble and roll kinematics."
    )
    r_t1.font.name = "Calibri"
    r_t1.font.size = Pt(9.2)
    r_t1.font.color.rgb = BODY_COLOR

    # Step 2
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(1)
    p_s2.paragraph_format.space_after = Pt(2.5)
    p_s2.paragraph_format.line_spacing = 1.12
    r_h2 = p_s2.add_run("• Step 2: Instant Pre-Alert & 15-Second Unconscious Failsafe: ")
    r_h2.font.name = "Calibri"
    r_h2.font.size = Pt(9.2)
    r_h2.font.bold = True
    r_h2.font.color.rgb = PRIMARY
    r_t2 = p_s2.add_run(
        "The millisecond an impact is detected, Rakshak-AI sends an ultra-lightweight 'suspected' ping to Police Command, giving patrol units a 15-second head start. "
        "Simultaneously, the phone sounds a 100 dB siren with an animated 15-second countdown. If the rider is uninjured, a single tap on 'I'M OK' cancels the alert. "
        "If unconscious, the timer expires and promotes the incident network-wide. If mobile data is dead, a local SQLite queue retries with exponential backoff "
        "and automatically triggers native dual-SIM GSM SMS alerts with GPS coordinates to family contacts."
    )
    r_t2.font.name = "Calibri"
    r_t2.font.size = Pt(9.2)
    r_t2.font.color.rgb = BODY_COLOR

    # Step 3
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_before = Pt(1)
    p_s3.paragraph_format.space_after = Pt(2.5)
    p_s3.paragraph_format.line_spacing = 1.12
    r_h3 = p_s3.add_run("• Step 3: C-V2X Autonomous Green Wave Signal Preemption: ")
    r_h3.font.name = "Calibri"
    r_h3.font.size = Pt(9.2)
    r_h3.font.bold = True
    r_h3.font.color.rgb = PRIMARY
    r_t3 = p_s3.add_run(
        "When an ambulance claims the case, its cockpit HUD projects a 950-meter forward radar geofence. Compatible with municipal traffic controllers "
        "(NTCIP 1202 / B-TRAC standards), oncoming signals transition conflicting phases through yellow clearance and turn SOLID GREEN 45 seconds before the ambulance arrives. "
        "CCTV Edge Vision (YOLOv8) monitors intersection congestion in real time, routing the ambulance via dynamic A* bypasses to avoid bottlenecks (saving 4.2 minutes in trials)."
    )
    r_t3.font.name = "Calibri"
    r_t3.font.size = Pt(9.2)
    r_t3.font.color.rgb = BODY_COLOR

    # Step 4
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_before = Pt(1)
    p_s4.paragraph_format.space_after = Pt(2.5)
    p_s4.paragraph_format.line_spacing = 1.12
    r_h4 = p_s4.add_run("• Step 4: Pre-Armed Trauma Care & Ayushman Bharat (ABHA) Health Card: ")
    r_h4.font.name = "Calibri"
    r_h4.font.size = Pt(9.2)
    r_h4.font.bold = True
    r_h4.font.color.rgb = PRIMARY
    r_t4 = p_s4.add_run(
        "Nearest trauma hospitals (e.g., Manipal, Cratis, Aster) receive dispatches via atomic PostgreSQL row locking to prevent conflicting claims. "
        "10 minutes prior to arrival, the trauma bay pre-arms Bay 01 and accesses the patient's verified ABHA record—revealing their confirmed blood group (B+) "
        "and critical contraindications (e.g., Penicillin anaphylaxis). Under strict Break-Glass access control, accessing medical records generates an immutable audit record, "
        "while police consoles are cryptographically barred from viewing clinical data."
    )
    r_t4.font.name = "Calibri"
    r_t4.font.size = Pt(9.2)
    r_t4.font.color.rgb = BODY_COLOR

    # -------------------------------------------------------------
    # SECTION 4: GROUNDED ENGINEERING & PRIVACY
    # -------------------------------------------------------------
    add_heading("4. Grounded Engineering Honesty & DPDP Act 2023 Compliance")
    
    add_para(
        "A common flaw in academic projects is claiming impossible metrics—like measuring 50G impacts or precise bone fractures from a phone in a pocket. "
        "Standard smartphone MEMS accelerometers physically saturate between 8G and 16G. In Rakshak-AI, I practice sensor honesty: if a reading exceeds hardware limits, "
        "it is explicitly flagged as 'Saturated' rather than displaying fabricated numbers. For highway crashes where phones may be destroyed, I developed a server-side "
        "'Trip Sentinel'. If a trip cruising above 60 km/h abruptly terminates heartbeats with zero deceleration signature, the server escalates through a 3-tier state machine to alert highway police."
    )

    add_para(
        "Regarding privacy, India's Digital Personal Data Protection (DPDP) Act 2023 requires strict purpose limitation. Rakshak-AI stores zero crash photos or clinical PII "
        "on persistent disk. Visual triage streams directly over WebRTC in volatile RAM with a 900-second TTL. Backend serializers actively sanitize database records "
        "before writing to PostgreSQL, keeping emergency health cards strictly in volatile Redis RAM with a 2-hour auto-purge."
    )

    # -------------------------------------------------------------
    # SECTION 5: WORKING IMPLEMENTATION & TESTING
    # -------------------------------------------------------------
    add_heading("5. Working Implementation & Verification")
    
    add_para(
        "Rakshak-AI is a fully functional, verified codebase with 93 passing automated backend tests: (1) 18/18 RBAC & Auth tests enforcing JWT cryptographic boundaries; "
        "(2) 29/29 Phase 17 Intelligence tests validating the suspected-to-promoted state machine, UUID idempotency deduplication, highway 60 km/h speed threshold, "
        "and blackspot spatial clustering; and (3) 46/46 Phase 18 Health Card tests validating ABDM adapter mock scenarios, Redis RAM storage, and break-glass audit logs. "
        "The React 18 dashboard compiles with zero errors (1,560 modules bundled), and the Flutter ambulance cockpit passes automated widget execution."
    )

    # -------------------------------------------------------------
    # SECTION 6: CONCLUSION & DEPLOYMENT FEASIBILITY
    # -------------------------------------------------------------
    add_heading("6. Feasibility & Deployment")
    
    add_para(
        "Because Rakshak-AI runs on everyday citizen smartphones and standard traffic controller protocols (NTCIP 1202), it requires zero bespoke hardware on vehicles "
        "and costs under ₹14,000 per intersection to deploy—compared to ₹8.5 Lakhs for proprietary roadside sensor grids. By replacing bystander panic, red light paralysis, "
        "and ER guesswork with automated orchestration, Rakshak-AI reclaims the Golden Hour, delivering an indigenous, reliable solution to protect Indian commuters."
    )

    # Save document
    doc.save(output_path)
    print(f"[SUCCESS] Concept Document created cleanly at: {output_path}")

if __name__ == "__main__":
    out1 = "CONCEPT_DOCUMENT_RAKSHAK_AI.docx"
    create_concept_document(out1)
    
    pitch_dir = os.path.join("Vikist Blueprint", "1. CONCEPT")
    if os.path.exists(pitch_dir):
        out2 = os.path.join(pitch_dir, "CONCEPT_DOCUMENT.docx")
        create_concept_document(out2)
